"""Parse raw RBI Excel releases (raw/rbi/...) into tidy tables in data/bank.db.

Tables written (all long format):
  sibc_credit       sectoral / industry-wise bank credit, Rs crore, one row per
                    (statement, code, as_of); newest release wins on overlap
  rates_group       WALR / term-deposit rates by bank group, monthly, %
  rates_sector      sector-wise WALR (fresh and outstanding), monthly, %
  rates_mclr_group  bank-group 1y MCLR min / max / median, monthly, %
  rates_mclr_bank   bank-wise 1y MCLR, monthly, %
  rates_benchmark   share of floating-rate loans by benchmark, quarterly, %
  cards             bank-wise ATM / PoS / QR / card stats, monthly; amounts are
                    converted from Rs '000 to Rs crore
  sop_scb           deposits, credit, G-sec investments of scheduled commercial
                    banks from the fortnightly Statement of Position PDFs, Rs crore
  data_notes        caveats copied from the RBI releases

Usage:  python src/parse_rbi.py
"""
import re
import sqlite3
import warnings
from datetime import datetime, date
from pathlib import Path

import openpyxl
import pandas as pd

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "raw" / "rbi"
DB = ROOT / "data" / "bank.db"

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def release_key(path):
    """'2026-August.xlsx' -> date(2026, 8, 1), used to rank releases."""
    m = re.match(r"(\d{4})-([A-Za-z]+)", path.stem)
    return date(int(m.group(1)), MONTHS[m.group(2)[:3].lower()], 1)


def to_month(v):
    """Excel cell (datetime or text like 'May-2016', 'Sept-2016') -> first of month."""
    if isinstance(v, datetime):
        return date(v.year, v.month, 1)
    if isinstance(v, date):
        return date(v.year, v.month, 1)
    if isinstance(v, str):
        m = re.match(r"\s*([A-Za-z]{3})[A-Za-z]*[-\s,]+(\d{4})\s*$", v)
        if m and m.group(1).lower() in MONTHS:
            return date(int(m.group(2)), MONTHS[m.group(1).lower()], 1)
    return None


def to_day(s):
    """SIBC header text like '31.Aug,2026' -> date."""
    m = re.fullmatch(r"\s*(\d{1,2})\.\s*([A-Za-z]{3})[A-Za-z]*\s*,\s*(\d{4})\s*", str(s))
    if not m:
        return None
    return date(int(m.group(3)), MONTHS[m.group(2).lower()], int(m.group(1)))


def data_month(d):
    """Month a SIBC reporting date stands for. Before Dec-2025 aggregates used the
    'last reporting fortnight', which often fell in the first week of the next
    month (e.g. 5-Sep-2025 is the August figure, 4-Apr-2025 the March year-end)."""
    if d.day <= 8:
        y, m = (d.year, d.month - 1) if d.month > 1 else (d.year - 1, 12)
        return date(y, m, 1)
    return date(d.year, d.month, 1)


def rows(path, sheet=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


# --------------------------------------------------------------------------- SIBC

LABEL_RE = re.compile(r"^\s*((?:[IVX]+|\d+(?:\.\d+)*)\.|\([ivx]+\))\s*(.+?)\s*$")


def parse_sibc_sheet(grid, statement, release):
    # Date header rows list the level columns (5 normally, 3 in the March
    # release) followed by growth columns; a new header row starts each block.
    out, dates = [], None
    for r in grid:
        parsed = []
        for c in r[1:]:
            d = to_day(c)
            if d is None:
                break
            parsed.append(d)
        if len(parsed) >= 3:
            dates = parsed
            continue
        label = r[0]
        if not isinstance(label, str) or dates is None:
            continue
        cells = r[1:1 + len(dates)]
        if not all(isinstance(c, (int, float)) for c in cells):
            continue
        text = " ".join(label.split())
        if text.startswith("Industries (2.1 to 2.19)"):
            code, name = "2", "Industry (total)"
        else:
            m = LABEL_RE.match(text)
            if not m:
                continue
            code, name = m.group(1).rstrip("."), m.group(2)
            if code.startswith("("):
                code = "PS." + code.strip("()")
        name = re.sub(r"\)\d+", ")", name)                         # '(NBFCs)2 of which,'
        name = re.sub(r",?\s*of which,?\s*$", "", name)
        name = re.sub(r"(?<=[A-Za-z\)])\d+$", "", name).strip()  # drop footnote digits
        for d, v in zip(dates, cells):
            out.append(dict(statement=statement, code=code, sector=name,
                            as_of=d, value_cr=float(v), release=release))
    return out


def build_sibc():
    recs = []
    for f in sorted((RAW / "sibc").glob("*.xlsx")):
        rel = release_key(f)
        recs += parse_sibc_sheet(rows(f, "Statement 1"), "sector", rel)
        recs += parse_sibc_sheet(rows(f, "Statement 2"), "industry", rel)
    df = pd.DataFrame(recs)
    df["as_of_month"] = df["as_of"].map(data_month)
    # newest release wins (revisions); keep the release it came from
    df = (df.sort_values("release")
            .drop_duplicates(["statement", "code", "as_of_month"], keep="last"))
    df["level"] = df["code"].map(lambda c: 0 if c in ("I", "II", "III") or c.startswith("PS.")
                                 else c.count(".") + 1)
    return df.sort_values(["statement", "code", "as_of"]).reset_index(drop=True)


# -------------------------------------------------------------------------- Rates

def find_header(grid, labels=("month", "month-end", "end month")):
    for i, r in enumerate(grid):
        if isinstance(r[1], str) and r[1].strip().lower() in labels:
            return i
    raise ValueError("header row not found")


def norm_group(s):
    return " ".join(str(s).split())


def parse_wide_monthly(grid, value_name):
    """Date in column 1, one series per column; header row labels the columns."""
    h = find_header(grid)
    heads = {j: norm_group(c) for j, c in enumerate(grid[h]) if j >= 2 and c}
    out = []
    for r in grid[h + 1:]:
        m = to_month(r[1])
        if m is None:
            continue
        for j, name in heads.items():
            v = r[j] if j < len(r) else None
            if isinstance(v, (int, float)):
                out.append({"month": m, value_name: name, "value": float(v)})
    return out


def parse_two_level(grid, sub_name):
    """Group header row followed by a sub-header row (e.g. Min/Max/Median)."""
    h = None
    for i, r in enumerate(grid):
        if isinstance(r[1], str) and r[1].strip().lower() in ("month", "end month"):
            h = i
            break
    top, sub = grid[h], grid[h + 1]
    # some sheets put the date label on the sub-header row instead
    if not any(isinstance(c, str) and c.strip() for c in sub[2:]):
        top, sub = grid[h - 1], grid[h]
    groups, cur = {}, None
    for j in range(2, len(sub)):
        if j < len(top) and top[j]:
            cur = norm_group(top[j])
        if sub[j]:
            groups[j] = (cur, norm_group(sub[j]))
    out = []
    for r in grid[h + 1:]:
        m = to_month(r[1])
        if m is None:
            continue
        for j, (g, s) in groups.items():
            v = r[j] if j < len(r) else None
            if isinstance(v, (int, float)) and s.lower() != "total":
                out.append({"month": m, "bank_group": g, sub_name: s, "value": float(v)})
    return out


def parse_bank_mclr(grid):
    h = next(i for i, r in enumerate(grid) if isinstance(r[2], str) and r[2].strip() == "Banks")
    months = {j: to_month(c) for j, c in enumerate(grid[h]) if j >= 3}
    out, group = [], None
    for r in grid[h + 1:]:
        if isinstance(r[1], str) and re.match(r"\(\w\)", r[1].strip()):
            group = norm_group(r[2])
            continue
        if not isinstance(r[2], str) or not isinstance(r[1], (int, float)):
            continue
        for j, m in months.items():
            v = r[j] if j < len(r) else None
            if m and isinstance(v, (int, float)):
                out.append({"bank": norm_group(r[2]), "bank_group": group,
                            "month": m, "value": float(v)})
    return out


RATE_SHEETS = {
    "Monthly Outstanding Loans WALR": "walr_outstanding",
    "Monthly Fresh Rupee Loans WALR": "walr_fresh",
    "Outstanding Rupee Term Deposits": "td_outstanding",
    "Fresh Rupee Term Deposits": "td_fresh",
}
SECTOR_SHEETS = {
    "Sector-wise Outstanding WALR": "walr_outstanding",
    "Sector-wise Fresh WALR": "walr_fresh",
}


def build_rates():
    latest = max((RAW / "rates").glob("*.xlsx"), key=release_key)
    wb_sheets = openpyxl.load_workbook(latest, read_only=True).sheetnames
    grp = []
    for sheet, series in RATE_SHEETS.items():
        for rec in parse_wide_monthly(rows(latest, sheet), "bank_group"):
            grp.append({"series": series, **rec})
    sec = []
    for sheet, series in SECTOR_SHEETS.items():
        if sheet in wb_sheets:
            for rec in parse_wide_monthly(rows(latest, sheet), "sector"):
                sec.append({"series": series, **rec})
    mclr_g = parse_two_level(rows(latest, "Bank Groupwise MCLR"), "stat")
    bench = parse_two_level(rows(latest, "Bank Groupwise Benchmark Share"), "benchmark")
    mclr_b = parse_bank_mclr(rows(latest, "Bank wise 1 Year MCLR"))
    return {
        "rates_group": pd.DataFrame(grp),
        "rates_sector": pd.DataFrame(sec),
        "rates_mclr_group": pd.DataFrame(mclr_g),
        "rates_benchmark": pd.DataFrame(bench),
        "rates_mclr_bank": pd.DataFrame(mclr_b),
    }, latest.name


# -------------------------------------------------------------------------- Cards

CARD_COLS = {  # column-number label in RBI's numbering row -> (metric, unit)
    1: ("atm_onsite", "count"), 2: ("atm_offsite", "count"), 3: ("pos", "count"),
    4: ("micro_atm", "count"), 5: ("bharat_qr", "count"), 6: ("upi_qr", "count"),
    7: ("credit_cards", "count"), 8: ("debit_cards", "count"),
    9: ("cc_pos_vol", "count"), 10: ("cc_pos_val", "rs_000"),
    11: ("cc_ecom_vol", "count"), 12: ("cc_ecom_val", "rs_000"),
    13: ("cc_other_vol", "count"), 14: ("cc_other_val", "rs_000"),
    15: ("cc_atm_cash_vol", "count"), 16: ("cc_atm_cash_val", "rs_000"),
    17: ("dc_pos_vol", "count"), 18: ("dc_pos_val", "rs_000"),
    19: ("dc_ecom_vol", "count"), 20: ("dc_ecom_val", "rs_000"),
    21: ("dc_other_vol", "count"), 22: ("dc_other_val", "rs_000"),
    23: ("dc_atm_cash_vol", "count"), 24: ("dc_atm_cash_val", "rs_000"),
    25: ("dc_pos_cash_vol", "count"), 26: ("dc_pos_cash_val", "rs_000"),
}
GROUPS = ("Public Sector Banks", "Private Sector Banks", "Foreign Banks",
          "Payment Banks", "Small Finance Banks")


def canon_bank(name):
    s = " ".join(str(name).upper().replace("&", " AND ").split())
    s = re.sub(r"[.,]", "", s)
    s = re.sub(r"\b(LTD|LIMITED|PLC|BSC|QPSC|CORP(ORATION)?)\b", "", s)
    return " ".join(s.split())


def parse_cards(path):
    grid = rows(path)
    title = next(str(r[1]) for r in grid if r[1] and "Month of" in str(r[1]))
    m = re.search(r"Month of\s+([A-Za-z]+)\s+(\d{4})", title)
    month = date(int(m.group(2)), MONTHS[m.group(1)[:3].lower()], 1)
    # locate RBI's column-numbering row (1, 2, ... 26) to pin column positions
    num_row = next(i for i, r in enumerate(grid)
                   if [c for c in r if isinstance(c, (int, float))][:3] == [1, 2, 3]
                   and len([c for c in r if isinstance(c, (int, float))]) == 26)
    colmap = {j: CARD_COLS[int(c)] for j, c in enumerate(grid[num_row])
              if isinstance(c, (int, float)) and int(c) in CARD_COLS}
    assert len(colmap) == 26, f"{path.name}: unexpected card layout"
    first_data_col = min(colmap)
    out, group = [], None
    for r in grid[num_row + 1:]:
        texts = [c for c in r[:first_data_col] if isinstance(c, str) and c.strip()]
        label = " ".join(texts[0].split()) if texts else ""
        if label in GROUPS:
            group = label
            continue
        if label.lower().startswith("note"):
            break
        is_total = label.lower() == "total"
        if not (group and (is_total or (len(texts) == 1 and any(
                isinstance(c, (int, float)) for c in r[:first_data_col])))):
            continue
        bank = "TOTAL" if is_total else canon_bank(label)
        for j, (metric, unit) in colmap.items():
            v = r[j] if j < len(r) else None
            if isinstance(v, (int, float)):
                val = float(v) / 1e4 if unit == "rs_000" else float(v)  # Rs'000 -> Rs crore
                out.append({"month": month, "bank_group": None if is_total else group,
                            "bank": bank, "bank_name_raw": label, "metric": metric,
                            "value": val, "unit": "rs_crore" if unit == "rs_000" else unit})
    return out


def build_cards():
    recs = []
    for f in sorted((RAW / "cards").glob("*.xlsx")):
        recs += parse_cards(f)
    return pd.DataFrame(recs)


# ---------------------------------------------------- Statement of Position (SoP)

SOP_ROWS = {  # label prefix on the PDF line -> metric
    "a) Deposits (other than from banks)": "deposits",
    "i) Demand": "deposits_demand",
    "ii) Time": "deposits_time",
    "V BALANCES WITH R.B.I.": "balances_with_rbi",
    "IV CASH": "cash",
    "a) Central & State Govt. securities": "investments_gsec",
    "VIII": "bank_credit",
}
NUM_RE = re.compile(r"^-?\d[\d,]*\.\d+")


def parse_sop(path):
    """Scheduled commercial banks columns (first three numbers on each line) of
    RBI's fortnightly 'Scheduled Banks' Statement of Position' PDF."""
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        words = pdf.pages[0].extract_words()
    # Cluster text words into lines (tops within 2pt), then attach each number to
    # the nearest line: numbers sit a fraction of a point off their label.
    text_w = sorted((w for w in words if not NUM_RE.match(w["text"])), key=lambda w: w["top"])
    lines = []
    for w in text_w:
        if lines and abs(w["top"] - lines[-1]["top"]) <= 2:
            lines[-1]["words"].append(w)
        else:
            lines.append({"top": w["top"], "words": [w], "nums": []})
    for w in words:
        if NUM_RE.match(w["text"]) and lines:
            ln = min(lines, key=lambda l: abs(l["top"] - w["top"]))
            if abs(ln["top"] - w["top"]) <= 4:
                ln["nums"].append(w)
    dates, out = None, []
    seen_liab_others = False
    for ln in lines:
        texts = [w["text"] for w in sorted(ln["words"], key=lambda w: w["x0"])]
        if dates is None:
            ds = [t.rstrip("#") for t in texts if re.fullmatch(r"\d{2}-[A-Za-z]{3}-\d{4}#?", t)]
            if len(ds) >= 3:
                dates = [datetime.strptime(d, "%d-%b-%Y").date() for d in ds[:3]]
            continue
        label = " ".join(texts)
        nums = [float(NUM_RE.match(w["text"]).group().replace(",", ""))
                for w in sorted(ln["nums"], key=lambda w: w["x0"])]
        if label.startswith("II LIABILITIES TO OTHERS"):
            seen_liab_others = True
        for prefix, metric in SOP_ROWS.items():
            # 'i) Demand' / 'ii) Time' occur only under deposits (item II)
            if label.startswith(prefix) and len(nums) >= 3:
                if metric.startswith("deposits_") and not seen_liab_others:
                    continue
                if any(r["metric"] == metric for r in out):
                    continue
                for d, v in zip(dates, nums[:3]):
                    out.append({"as_of": d, "as_of_month": data_month(d), "metric": metric,
                                "value_cr": v, "release": path.stem})
    return out


def build_sop():
    recs = []
    for f in sorted((RAW / "sop").glob("*.pdf")):
        recs += parse_sop(f)
    df = pd.DataFrame(recs)
    if df.empty:
        return df
    # keep month-end figures; for each (metric, month) prefer the newest release
    df["release_date"] = df["release"].map(lambda s: datetime.strptime(s, "%Y-%B-%d").date())
    df = df[df["as_of"] != df["release_date"].map(lambda d: date(d.year, d.month, 15))]
    df = (df.sort_values(["release_date", "as_of"])
            .drop_duplicates(["metric", "as_of_month"], keep="last")
            .drop(columns="release_date"))
    return df.sort_values(["metric", "as_of_month"]).reset_index(drop=True)


# -------------------------------------------------------------------------- Notes

NOTES = [
    ("sibc", "Aggregate bank/food/non-food credit come from the fortnightly Section-42 return (all SCBs); "
             "sector/industry splits come from the SIBC return (select banks, ~95% of non-food credit). "
             "Sector totals therefore do not add exactly to non-food credit."),
    ("sibc", "From 31-Dec-2025 the reporting date is month-end (Banking Laws (Amendment) Act 2025); earlier "
             "values are the last reporting fortnight, so YoY growth across that date mixes the two bases."),
    ("sibc", "Since May-2024 one bank reclassified some agricultural loans as 'Loans against gold jewellery', "
             "inflating gold-loan growth and depressing agri growth for the following 12 months."),
    ("rates", "Data from July 2023 include the impact of the HDFC Ltd merger with HDFC Bank."),
    ("cards", "Transaction values converted from Rs '000 to Rs crore; counts are actual numbers."),
    ("sop", "Scheduled commercial banks incl. RRBs, SFBs and payment banks. Year-ago columns are the "
            "corresponding reporting Friday under the old fortnight definition (until Nov-2026)."),
]


def main():
    DB.parent.mkdir(exist_ok=True)
    sibc = build_sibc()
    rates, rates_file = build_rates()
    cards = build_cards()
    sop = build_sop()
    with sqlite3.connect(DB) as con:
        sop.to_sql("sop_scb", con, if_exists="replace", index=False)
        sibc.to_sql("sibc_credit", con, if_exists="replace", index=False)
        for name, df in rates.items():
            df.to_sql(name, con, if_exists="replace", index=False)
        cards.to_sql("cards", con, if_exists="replace", index=False)
        pd.DataFrame(NOTES, columns=["dataset", "note"]).to_sql(
            "data_notes", con, if_exists="replace", index=False)
    print(f"sibc_credit: {len(sibc):6d} rows  {sibc.as_of.min()} .. {sibc.as_of.max()}")
    print(f"rates (from {rates_file}):")
    for n, df in rates.items():
        print(f"  {n:17s} {len(df):6d} rows  {df.month.min()} .. {df.month.max()}")
    print(f"cards:       {len(cards):6d} rows  {cards.month.min()} .. {cards.month.max()}"
          f"  banks={cards.bank.nunique()}")
    if not sop.empty:
        print(f"sop_scb:     {len(sop):6d} rows  {sop.as_of.min()} .. {sop.as_of.max()}"
              f"  metrics={sorted(sop.metric.unique())}")
    print(f"-> {DB}")


if __name__ == "__main__":
    main()
